import csv
import io
import json

from openpyxl import load_workbook

import app as flask_app


def _write_dataset(path, name="users", unit="people", q1="10", q2="12"):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["metric", "q1", "q2", "unit"])
        writer.writerow([name, q1, q2, unit])


def test_dashboard_api_and_exports(tmp_path):
    data = tmp_path / "metrics.csv"
    _write_dataset(data)
    flask_app.app.config.update(TESTING=True, DATASET_PATH=data)
    client = flask_app.app.test_client()

    page = client.get("/")
    assert page.status_code == 200
    assert b"NovaFlow" in page.data
    assert b"users" in page.data
    assert b"only when you press the button" in page.data

    response = client.get("/api/summary")
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["summary"]["total_metrics"] == 1
    assert payload["metrics"][0]["percentage_change"] == 20

    exported_json = client.get("/export/json")
    assert exported_json.status_code == 200
    assert exported_json.mimetype == "application/json"
    assert exported_json.headers["Content-Disposition"].startswith("attachment;")
    assert json.loads(exported_json.data)["metrics"][0]["name"] == "users"

    exported_excel = client.get("/export/excel")
    assert exported_excel.status_code == 200
    assert exported_excel.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert exported_excel.headers["Content-Disposition"].startswith("attachment;")
    workbook = load_workbook(filename=io.BytesIO(exported_excel.data), read_only=True)
    sheet = workbook.active
    assert sheet["A2"].value == "users"
    assert sheet["B2"].data_type == "n"


def test_excel_export_neutralizes_formula_prefixes_in_names_and_units(tmp_path):
    prefixes = ["=", "+", "-", "@", "\t=", "\r=", " =1+2", "\t@SUM(A1:A2)"]
    data = tmp_path / "formula.csv"
    _write_dataset(data, name="ordinary", unit="ordinary")
    with data.open("a", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for index, prefix in enumerate(prefixes):
            writer.writerow([prefix + "payload", index + 1, index + 2, prefix + "unit"])

    flask_app.app.config.update(TESTING=True, DATASET_PATH=data)
    response = flask_app.app.test_client().get("/export/excel")
    assert response.status_code == 200
    workbook = load_workbook(filename=io.BytesIO(response.data), read_only=True)
    sheet = workbook.active
    assert sheet["A2"].value == "ordinary"
    assert sheet["F2"].value == "ordinary"
    for row_number in range(3, 3 + len(prefixes)):
        for column in ("A", "F"):
            cell = sheet[f"{column}{row_number}"]
            assert cell.value.startswith("'")
            assert cell.data_type == "s"
        assert sheet[f"B{row_number}"].data_type == "n"
        assert sheet[f"C{row_number}"].data_type == "n"


def test_invalid_and_missing_dataset_return_actionable_errors(tmp_path):
    flask_app.app.config.update(TESTING=True, DATASET_PATH=tmp_path / "absent.csv")
    client = flask_app.app.test_client()
    assert client.get("/").status_code == 404
    assert client.get("/health").get_json() == {"status": "ok", "service": "NovaFlow"}

    bad = tmp_path / "bad.csv"
    bad.write_text("metric,q1,q2,unit\nusers,nan,3,people\n", encoding="utf-8")
    flask_app.app.config["DATASET_PATH"] = bad
    response = client.get("/api/summary")
    assert response.status_code == 400
    assert "finite number" in response.get_json()["error"]


def test_ai_endpoint_requires_opt_in_and_explicit_confirmation(tmp_path, monkeypatch):
    data = tmp_path / "metrics.csv"
    _write_dataset(data)
    flask_app.app.config.update(TESTING=True, DATASET_PATH=data)
    client = flask_app.app.test_client()
    monkeypatch.delenv("NOVAFLOW_AI_ENABLED", raising=False)
    monkeypatch.setattr(flask_app, "generate_ai_insight", lambda metrics: (_ for _ in ()).throw(AssertionError("must not call AI")))

    disabled = client.post("/api/ai-insights", json={"confirm": True})
    assert disabled.status_code == 503
    assert "disabled" in disabled.get_json()["error"]

    monkeypatch.setenv("NOVAFLOW_AI_ENABLED", "true")
    unconfirmed = client.post("/api/ai-insights", json={"confirm": False})
    assert unconfirmed.status_code == 400


def test_ai_endpoint_returns_labeled_mocked_insight(tmp_path, monkeypatch):
    data = tmp_path / "metrics.csv"
    _write_dataset(data)
    flask_app.app.config.update(TESTING=True, DATASET_PATH=data)
    monkeypatch.setenv("NOVAFLOW_AI_ENABLED", "true")
    monkeypatch.setattr(flask_app, "generate_ai_insight", lambda metrics: "Users increased by 20%.")
    result = flask_app.app.test_client().post("/api/ai-insights", json={"confirm": True})
    assert result.status_code == 200
    assert result.get_json()["generated_by"] == "AI"
    assert result.get_json()["insight"] == "Users increased by 20%."
    assert "verify" in result.get_json()["notice"]
